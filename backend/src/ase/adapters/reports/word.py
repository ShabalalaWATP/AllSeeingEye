"""Professional A4 Word rendering for the bounded report document contract."""

from __future__ import annotations

import io
import re
from typing import Any
from urllib.parse import urlsplit

from docx import Document
from docx.document import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

from ase.adapters.reports.document_sections import leading_identity
from ase.adapters.reports.document_validation import validate_document_content
from ase.adapters.reports.figure_validation import VerifiedFigure, verify_document_figures
from ase.adapters.reports.word_styles import (
    add_masthead,
    add_page_furniture,
    add_review_band,
    add_section_number,
    configure_styles,
)
from ase.domain.report_documents import (
    BlockKind,
    DocumentFigure,
    DocumentInline,
    DocumentReference,
    DocumentTable,
    ReportDocument,
)

_CITATION_NUMBER = re.compile(r"\d+")


def _safe_url(value: str | None) -> str | None:
    if not value or len(value) > 2_048 or any(ord(char) < 32 for char in value):
        return None
    parsed = urlsplit(value)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        return None
    if parsed.username is not None or parsed.password is not None:
        return None
    return value


def _hyperlink_run(text: str) -> Any:
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), "16748A")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((colour, underline))
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.extend((properties, node))
    return run


def _add_internal_link(paragraph: Paragraph, text: str, anchor: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.append(_hyperlink_run(text))
    paragraph._p.append(hyperlink)


def _add_external_link(paragraph: Paragraph, text: str, url: str) -> None:
    relationship = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    hyperlink.append(_hyperlink_run(text))
    paragraph._p.append(hyperlink)


def _add_bookmark(paragraph: Paragraph, name: str, identifier: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(identifier))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(identifier))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _add_citation_inline(paragraph: Paragraph, inline: DocumentInline) -> None:
    matches = tuple(_CITATION_NUMBER.finditer(inline.text))
    if tuple(int(match.group()) for match in matches) != inline.citation_numbers:
        paragraph.add_run(inline.text)
        return
    numbers = iter(inline.citation_numbers)
    cursor = 0
    for match in matches:
        try:
            number = next(numbers)
        except StopIteration:
            break
        paragraph.add_run(inline.text[cursor : match.start()])
        _add_internal_link(paragraph, match.group(), f"reference-{number}")
        cursor = match.end()
    paragraph.add_run(inline.text[cursor:])


def _add_inlines(paragraph: Paragraph, text: str, inlines: tuple[DocumentInline, ...]) -> None:
    if not inlines:
        paragraph.add_run(text)
        return
    for inline in inlines:
        if inline.citation_numbers:
            _add_citation_inline(paragraph, inline)
        else:
            paragraph.add_run(inline.text)


def _add_reference(paragraph: Paragraph, reference: DocumentReference, text: str) -> None:
    _add_bookmark(paragraph, f"reference-{reference.number}", reference.number)
    urls = tuple(
        dict.fromkeys(
            url
            for candidate in (reference.url, reference.archive_url)
            if (url := _safe_url(candidate)) is not None
        )
    )
    cursor = 0
    for url in urls:
        start = text.find(url, cursor)
        if start < 0:
            continue
        paragraph.add_run(text[cursor:start])
        _add_external_link(paragraph, url, url)
        cursor = start + len(url)
    paragraph.add_run(text[cursor:])


def _add_table(word: WordDocument, table: DocumentTable) -> None:
    title = word.add_paragraph(table.title, style="Heading 2")
    title.paragraph_format.keep_with_next = True
    native = word.add_table(rows=1, cols=len(table.columns))
    native.style = "Light List Accent 1"
    native.autofit = True
    header = native.rows[0]
    header_properties = header._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_properties.append(repeat)
    for cell, value in zip(header.cells, table.columns, strict=True):
        cell.text = value
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(3)
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(8.5)
    for row in table.rows:
        cells = native.add_row().cells
        for output, cell_value in zip(cells, row, strict=True):
            output.text = ""
            body = output.paragraphs[0]
            body.paragraph_format.space_after = Pt(3)
            _add_inlines(body, cell_value.text, cell_value.inlines)
            for run in body.runs:
                run.font.size = Pt(9)
    if table.caption:
        caption = word.add_paragraph(table.caption, style="ASE Metadata")
        caption.paragraph_format.space_after = Pt(10)


def _add_figure(word: WordDocument, figure: DocumentFigure, verified: VerifiedFigure) -> None:
    paragraph = word.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    max_width, max_height = 6.7, 6.3
    scale = min(max_width / figure.width_px, max_height / figure.height_px)
    shape = paragraph.add_run().add_picture(
        io.BytesIO(verified.content),
        width=Inches(figure.width_px * scale),
        height=Inches(figure.height_px * scale),
    )
    shape._inline.docPr.set("descr", figure.alt_text)
    shape._inline.docPr.set("title", figure.title)
    caption = word.add_paragraph(style="Caption")
    caption.paragraph_format.keep_with_next = False
    caption.add_run(f"{figure.title}. {figure.caption}")
    if figure.citation_numbers:
        citation = f" [{', '.join(str(number) for number in figure.citation_numbers)}]"
        _add_citation_inline(
            caption,
            DocumentInline(citation, "ltr", figure.citation_numbers),
        )


def render_docx(document: ReportDocument) -> bytes:
    validate_document_content(document)
    figures = verify_document_figures(document)
    word = Document()
    section = word.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin, section.bottom_margin = Inches(0.85), Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.78)
    styles = configure_styles(word)
    head, remaining = leading_identity(document)
    add_masthead(
        word, head, lambda paragraph, block: _add_inlines(paragraph, block.text, block.inlines)
    )
    reference_index = 0
    section_number = 0
    for block in remaining:
        if block.kind is BlockKind.TABLE and block.table:
            _add_table(word, block.table)
            continue
        if block.kind is BlockKind.FIGURE and block.figure:
            _add_figure(word, block.figure, figures[id(block.figure)])
            continue
        if block.kind is BlockKind.LIST:
            style = "List Number" if block.ordered else "List Bullet"
            for item in block.items:
                paragraph = word.add_paragraph(style=style)
                _add_inlines(paragraph, item.text, item.inlines)
            continue
        paragraph = word.add_paragraph(style=styles[block.kind])
        paragraph.paragraph_format.widow_control = True
        if block.kind in {BlockKind.HEADING, BlockKind.ANNEX}:
            section_number += 1
            add_section_number(paragraph, section_number)
        if block.kind is BlockKind.REFERENCE and reference_index < len(document.references):
            _add_reference(paragraph, document.references[reference_index], block.text)
            reference_index += 1
        else:
            _add_inlines(paragraph, block.text, block.inlines)
        if block.kind is BlockKind.WARNING:
            add_review_band(paragraph)
        if block.kind is BlockKind.ANNEX:
            paragraph.paragraph_format.page_break_before = True
    add_page_furniture(word, document)
    word.core_properties.title = document.title
    word.core_properties.author = "The All Seeing Eye"
    word.core_properties.subject = document.reference
    word.core_properties.comments = ""
    stream = io.BytesIO()
    word.save(stream)
    return stream.getvalue()
