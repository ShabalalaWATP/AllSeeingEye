"""Read exported disclosure back while keeping frozen research and sources unchanged."""

import io
import json
import zipfile
from dataclasses import replace
from uuid import uuid4

import pytest
from defusedxml.ElementTree import fromstring
from docx import Document
from pypdf import PdfReader

from ase.adapters.reports.evidence_package import FrozenEvidencePackageRenderer, evidence_geojson
from ase.adapters.reports.pdf import render_pdf
from ase.adapters.reports.print_html import render_print_html
from ase.adapters.reports.word import render_docx
from ase.application.reports.ai_disclosure import report_ai_disclosure
from ase.application.reports.document import build_document
from ase.application.reports.drafting import no_evidence_draft
from ase.domain.ai_disclosure import AI_LABEL, AiDisclosure
from ase.domain.errors import InvalidRequest
from ase.domain.evidence_attributes import EvidenceAttribute
from ase.domain.llm import LlmProvider, LlmRole
from ase.domain.model_routing import ModelRoutingRecord, RoutedModel
from ase.domain.report_records import body_to_dict, evidence_to_list
from ase.domain.reports import ReportStatus
from report_documents_helpers import document_records


def _routed(provider=LlmProvider.BEDROCK):
    record, version = document_records()
    profile_id = uuid4()
    route = RoutedModel(
        LlmRole.ASSESSMENT,
        profile_id,
        1,
        "configured-alias",
        None,
        1024,
        0.2,
        version.created_at,
        provider,
    )
    version = replace(
        version,
        profile_id=profile_id,
        model="recorded-response-model",
        model_routing=ModelRoutingRecord("legacy", None, None, (route,)),
    )
    return record, version


def test_pdf_disclosure_survives_readback_on_title_and_every_page():
    record, version = _routed()
    document = build_document(record, version)
    output = render_pdf(document)
    pdf = PdfReader(io.BytesIO(output))
    pages = [page.extract_text() for page in pdf.pages]
    assert len(pages) > 1
    assert pages[0].count(AI_LABEL) >= 2
    assert all(AI_LABEL in page for page in pages)
    assert all("recorded-response-model" in page and "Amazon Bedrock" in page for page in pages)
    assert pdf.metadata["/ase:ai-generated"] == "true"
    assert pdf.metadata["/ase:ai-model"] == "recorded-response-model"
    assert pdf.metadata["/ase:ai-provider"] == "bedrock"
    assert "AI-generated" in pdf.metadata.keywords
    assert pdf.metadata.title == record.title
    assert pdf.outline
    assert output == render_pdf(document)


def test_docx_has_visible_core_and_typed_custom_metadata():
    record, version = _routed()
    content = render_docx(build_document(record, version))
    word = Document(io.BytesIO(content))
    assert any(AI_LABEL in p.text and version.model in p.text for p in word.paragraphs)
    assert "AI-generated" in word.core_properties.keywords
    assert version.model in word.core_properties.comments
    for footer in (word.sections[0].footer, word.sections[0].first_page_footer):
        assert any(AI_LABEL in p.text and version.model in p.text for p in footer.paragraphs)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        root = fromstring(archive.read("docProps/custom.xml"))
        props = {prop.attrib["name"]: next(iter(prop)) for prop in root}
        assert props["ase:ai-generated"].tag.endswith("}bool")
        assert props["ase:ai-generated"].text == "true"
        assert props["ase:ai-model"].text == version.model
        assert props["ase:ai-provider"].text == "bedrock"
        assert "custom-properties" in archive.read("_rels/.rels").decode()
    assert content == render_docx(build_document(record, version))


def test_provenance_is_selected_frozen_version_not_latest_or_scope_fields():
    record, version = _routed()
    record.latest_version = 99
    record.scope = {"model": "client-forgery", "provider": "client-forgery"}
    document = build_document(record, version)
    assert document.ai_disclosure == AiDisclosure("recorded-response-model", "bedrock")
    assert "client-forgery" not in document.ai_disclosure.statement
    assert version.model_routing.profiles[0].model == "configured-alias"


def test_missing_and_mismatched_provider_provenance_stays_unknown():
    record, version = document_records()
    assert report_ai_disclosure(version).provider == "unknown"
    _, routed = _routed()
    assert report_ai_disclosure(replace(routed, profile_id=uuid4())).provider == "unknown"
    assert "Provider: Not recorded" in report_ai_disclosure(version).statement
    assert record.id == version.report_id


def test_openai_compatible_is_not_claimed_to_be_hosted_by_openai():
    _, version = _routed(LlmProvider.OPENAI_COMPATIBLE)
    disclosure = report_ai_disclosure(version)
    assert disclosure.provider == "openai_compatible"
    assert "host not recorded" in disclosure.statement


@pytest.mark.parametrize("model", ["bad\x00model", "bad\u202emodel", "x" * 2049])
def test_unsafe_saved_model_identifiers_fail_safely(model):
    _, version = document_records()
    with pytest.raises(InvalidRequest, match="saved model provenance"):
        report_ai_disclosure(replace(version, model=model))


def test_empty_legacy_model_remains_explicitly_unknown():
    _, version = document_records()
    disclosure = report_ai_disclosure(replace(version, model=""))
    assert "Model: Not recorded" in disclosure.statement
    assert "Model: Not recorded" in disclosure.footer


def test_long_literal_provenance_is_bounded_and_escaped_in_all_renderers():
    record, version = _routed()
    model = "<script>&" + "W" * 2039
    document = build_document(record, replace(version, model=model))
    pdf = PdfReader(io.BytesIO(render_pdf(document)))
    assert pdf.metadata["/ase:ai-model"] == model
    assert "full model at title" in pdf.pages[-1].extract_text()
    word = Document(io.BytesIO(render_docx(document)))
    assert any(model in p.text for p in word.paragraphs)
    html = render_print_html(document).decode()
    assert "&lt;script&gt;&amp;" in html
    assert "<script>" not in html
    assert 'class="ai-footer"' in html and "position: fixed" in html


def test_projection_does_not_change_body_blocks_evidence_or_saved_markdown():
    record, version = _routed()
    before = (body_to_dict(version.body), evidence_to_list(version.evidence), version.markdown)
    document = build_document(record, version)
    assert all(AI_LABEL not in block.text for block in document.blocks)
    package = FrozenEvidencePackageRenderer().render(record, version)
    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        assert archive.read("report.md").decode() == before[2]
        assert json.loads(archive.read("report.json"))["body"] == json.loads(json.dumps(before[0]))
        assert json.loads(archive.read("evidence.json")) == json.loads(json.dumps(before[1]))
        geo = json.loads(archive.read("evidence.geojson"))
    assert geo["metadata"]["report_ai_disclosure"]["ase:ai-model"] == version.model
    assert geo["metadata"]["application_ai_generated_fields"] == []
    assert geo["metadata"]["source_authorship"] == "unknown"
    assert geo["features"] == evidence_geojson(replace(version, model="different"))["features"]
    assert before == (
        body_to_dict(version.body),
        evidence_to_list(version.evidence),
        version.markdown,
    )


def test_non_ai_documents_do_not_acquire_generation_markers():
    record, version = document_records()
    document = replace(build_document(record, version), ai_disclosure=None)
    pdf = PdfReader(io.BytesIO(render_pdf(document)))
    assert "/ase:ai-generated" not in pdf.metadata
    with zipfile.ZipFile(io.BytesIO(render_docx(document))) as archive:
        assert "docProps/custom.xml" not in archive.namelist()
    assert 'class="ai-footer"' not in render_print_html(document).decode()


def test_unicode_model_provenance_is_recoverable_without_missing_pdf_glyphs():
    record, version = _routed()
    model = "模型" * 100
    content = render_pdf(build_document(record, replace(version, model=model)))
    pdf = PdfReader(io.BytesIO(content))
    assert pdf.metadata["/ase:ai-model"] == model
    first = pdf.pages[0].extract_text()
    assert "[U+6A21][U+578B]" in first
    assert "full identifier at title" in first
    assert "[U+0000]" not in first


def test_geojson_marks_recorded_ai_visual_hypothesis_titles_without_altering_them():
    _, version = document_records()
    title = "Unverified candidate: model-proposed location"
    hypothesis = replace(
        version.evidence[0],
        source_id="research_media",
        title=title,
        attributes=(EvidenceAttribute("provenance_status", "unverified AI visual hypothesis"),),
    )
    geo = evidence_geojson(replace(version, evidence=(hypothesis, *version.evidence[1:])))
    assert geo["metadata"]["application_ai_generated_fields"] == ["/features/0/properties/title"]
    assert "other fields unclassified" in geo["metadata"]["ai_field_classification"]
    assert geo["features"][0]["properties"]["title"] == title
    fixed_heading = evidence_geojson(
        replace(version, evidence=(replace(hypothesis, title="Assessment"),))
    )
    assert fixed_heading["metadata"]["application_ai_generated_fields"] == []


def test_no_evidence_gap_notice_is_not_attributed_to_the_configured_model():
    record, version = _routed()
    draft = no_evidence_draft("configured-but-not-called")
    version = replace(
        version,
        body=draft.body,
        findings=tuple(draft.findings),
        model=draft.model,
        attempts=draft.attempts,
        evidence=(),
    )
    document = build_document(record, version)
    pdf = PdfReader(io.BytesIO(render_pdf(document)))
    assert pdf.metadata["/ase:ai-generated"] == "false"
    assert pdf.metadata["/ase:ai-model"] == ""
    assert "AI-generated" not in pdf.metadata.keywords
    assert "No AI assessment was generated" in pdf.pages[0].extract_text()
    content = render_docx(document)
    word = Document(io.BytesIO(content))
    assert any("No AI assessment was generated" in p.text for p in word.paragraphs)
    assert "AI-generated" not in word.core_properties.keywords
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        root = fromstring(archive.read("docProps/custom.xml"))
        props = {p.attrib["name"]: next(iter(p)).text for p in root}
        assert props["ase:ai-generated"] == "false"
    assert (
        evidence_geojson(version)["metadata"]["report_ai_disclosure"]["ase:ai-generated"] == "false"
    )


def test_resumed_zero_attempt_or_failed_generated_body_keeps_ai_disclosure():
    record, version = _routed()
    resumed = replace(version, attempts=0, status=ReportStatus.FAILED)
    assert build_document(record, resumed).ai_disclosure.assessment_generated is True
