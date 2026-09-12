"""The reader and downloads share one clean, cited report product."""

import io
import zipfile
from dataclasses import replace

import pytest
from docx import Document
from pypdf import PdfReader

from ase.adapters.reports.documents import ReportDocumentRenderer
from ase.adapters.reports.html_document import render_html
from ase.api.schemas_reports import ReportOut
from ase.application.reports.document import build_document
from ase.application.reports.publication_markdown import render_document_markdown
from ase.domain.report_documents import (
    BlockKind,
    DocumentBlock,
    DocumentFigure,
    DocumentInline,
    DocumentReference,
    ExportFormat,
    ReportDocument,
)
from ase.domain.reports import ChangeFromPrevious, WatchCondition
from report_documents_helpers import document_records


def test_publication_is_reader_facing_and_references_follow_first_citation() -> None:
    record, version = document_records()
    document = build_document(record, version)
    text = "\n".join(block.text for block in document.blocks)

    assert [reference.evidence_label for reference in document.references] == ["E1", "E2", "E3"]
    assert [reference.number for reference in document.references] == [1, 2, 3]
    assert "Executive summary" in text
    assert "Findings and analysis" in text
    assert "Limitations and confidence" in text
    assert "References" in text
    assert any(block.kind is BlockKind.TABLE for block in document.blocks)
    assert "[1]" in text and "[E1" not in text
    for hidden in (
        version.model,
        "Attempts:",
        "Validator findings",
        "Frozen evidence annex",
        "Search terms:",
        version.evidence[0].content_hash,
    ):
        assert hidden not in text


def test_publication_markdown_uses_same_sections_table_and_linked_references() -> None:
    record, version = document_records()
    document = build_document(record, version)
    markdown = render_document_markdown(document)

    assert markdown.startswith(f"# {record.title}\n")
    assert "## Executive summary" in markdown
    assert "| Date | Source | Reported item |" in markdown
    assert "[1](#reference-1)" in markdown
    assert "### Reference 1" in markdown
    assert "[Open source](https://example.org/report)" in markdown
    assert (
        "[Archived copy](https://web.archive.org/web/20260901/https://example.org/report)"
        in markdown
    )
    assert "## Evidence annex" not in markdown
    assert version.model not in markdown


def test_publication_preserves_contrary_reporting_change_and_warning_fields() -> None:
    record, version = document_records()
    judgement = replace(
        version.body.key_judgements[0],
        contradicting_evidence=("E3",),
        change_from_previous=ChangeFromPrevious.STRENGTHENED,
        indicators=("Additional armoured columns arrive",),
    )
    body = replace(
        version.body,
        key_judgements=(judgement,),
        reporting=(),
        assessment=(),
        alternative_hypotheses=(),
        indicators_and_warning=replace(
            version.body.indicators_and_warning,
            watch_condition=WatchCondition.CRITICAL,
            changes=(),
        ),
        collection_recommendations=(),
    )
    document = build_document(record, replace(version, body=body))
    renderer = ReportDocumentRenderer()
    word = Document(io.BytesIO(renderer.render(document, ExportFormat.DOCX)))
    outputs = (
        "\n".join(block.text for block in document.blocks),
        render_document_markdown(document),
        "\n".join(paragraph.text for paragraph in word.paragraphs),
        "\n".join(
            page.extract_text()
            for page in PdfReader(io.BytesIO(renderer.render(document, ExportFormat.PDF))).pages
        ),
    )

    for raw_output in outputs:
        output = " ".join(raw_output.split())
        assert "Contrary reporting" in output
        assert "Change from previous assessment: strengthened" in output
        assert "Additional armoured columns arrive" in output
        assert "Watch condition: Critical" in output
    assert [reference.evidence_label for reference in document.references] == ["E1", "E2", "E3"]


def test_word_and_pdf_references_retain_source_and_archived_copy() -> None:
    record, version = document_records()
    document = build_document(record, version)
    renderer = ReportDocumentRenderer()
    word = renderer.render(document, ExportFormat.DOCX)
    with zipfile.ZipFile(io.BytesIO(word)) as archive:
        relationships = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.endswith(".rels")
        )
    assert 'Target="https://example.org/report"' in relationships
    assert (
        'Target="https://web.archive.org/web/20260901/https://example.org/report"' in relationships
    )
    pdf = renderer.render(document, ExportFormat.PDF)
    reader = PdfReader(io.BytesIO(pdf))
    pdf_text = "\n".join(page.extract_text() for page in reader.pages)
    assert "Source: https://example.org/report" in pdf_text
    assert "Archived copy:" in pdf_text
    assert "web.archive.org/web/20260901/https://example.org/report" in pdf_text
    assert f"Accessed {document.references[0].accessed_at}" in pdf_text
    targets = {
        action.get("/URI")
        for page in reader.pages
        for annotation in page.get("/Annots", ())
        if (action := annotation.get_object().get("/A")) is not None
    }
    assert "https://example.org/report" in targets
    assert "https://web.archive.org/web/20260901/https://example.org/report" in targets
    html = render_html(document).decode()
    assert 'href="https://example.org/report"' in html
    assert 'href="https://web.archive.org/web/20260901/https://example.org/report"' in html
    assert f"Accessed {document.references[0].accessed_at}" in html


def test_report_response_contains_the_exact_publication_contract() -> None:
    record, version = document_records()
    payload = ReportOut.build(record, version).model_dump(mode="json")
    publication = payload["version"]["publication"]

    assert publication["schema_version"] == 1
    assert publication["reference"].endswith("version 1")
    assert publication["references"][0]["number"] == 1
    table = next(block["table"] for block in publication["blocks"] if block["table"])
    assert table["columns"] == ["Date", "Source", "Reported item"]
    assert table["rows"][0][2]["inlines"][-1]["citation_numbers"]


def test_failed_publication_is_explicit_without_an_empty_reference_annex() -> None:
    record, version = document_records()
    failed = replace(
        version,
        status=version.status.FAILED,
        body=replace(
            version.body,
            key_judgements=(),
            reporting=(),
            assessment=(),
            alternative_hypotheses=(),
        ),
        evidence=(),
        markdown="",
    )
    document = build_document(record, failed)
    text = "\n".join(block.text for block in document.blocks)

    assert "FAILED GENERATION" in text
    assert "No assessed conclusion was produced" in text
    assert not document.references
    assert "References" not in text


def test_document_rejects_orphan_inline_and_figure_citations() -> None:
    reference = DocumentReference(1, "E1", "Title", "Publisher", None, "2026-09-12")
    with pytest.raises(ValueError, match="orphan citation"):
        ReportDocument(
            "Title",
            "Ref",
            (
                DocumentBlock(
                    BlockKind.TEXT,
                    "[2]",
                    (DocumentInline("[2]", citation_numbers=(2,)),),
                ),
            ),
            references=(reference,),
        )
    figure = DocumentFigure(
        "Map",
        "Caption",
        "Accessible description",
        b"not-decoded-at-domain-boundary",
        "image/png",
        10,
        10,
        (2,),
    )
    with pytest.raises(ValueError, match="orphan citation"):
        ReportDocument(
            "Title",
            "Ref",
            (DocumentBlock(BlockKind.FIGURE, "Caption", figure=figure),),
            references=(reference,),
        )
