"""Native report elements retain safe structure across document renderers."""

import io
import zipfile

import pytest
from docx import Document
from PIL import Image as PillowImage
from pypdf import PdfReader

from ase.adapters.reports import figure_validation
from ase.adapters.reports.html_document import render_html
from ase.adapters.reports.pdf import render_pdf
from ase.adapters.reports.word import render_docx
from ase.domain.errors import InvalidRequest
from ase.domain.report_documents import (
    BlockKind,
    DocumentBlock,
    DocumentFigure,
    DocumentInline,
    DocumentListItem,
    DocumentReference,
    DocumentTable,
    DocumentTableCell,
    ReportDocument,
)


def _png() -> bytes:
    stream = io.BytesIO()
    PillowImage.new("RGB", (320, 160), "#16748A").save(stream, format="PNG")
    return stream.getvalue()


def structured_document(*, figure_content: bytes | None = None) -> ReportDocument:
    citation = DocumentInline("[1]", "ltr", (1,))
    figure = DocumentFigure(
        "Figure 1",
        "Locations reported by the retained source.",
        "Two reported locations on a blue background.",
        figure_content or _png(),
        "image/png",
        320,
        160,
        (1,),
    )
    return ReportDocument(
        "Professional report",
        "Report fixture | version 1",
        (
            DocumentBlock(BlockKind.TITLE, "Professional report"),
            DocumentBlock(
                BlockKind.TEXT,
                "A source-backed finding [1]",
                (DocumentInline("A source-backed finding "), citation),
            ),
            DocumentBlock(
                BlockKind.LIST,
                "First reported point [1]",
                items=(
                    DocumentListItem(
                        "First reported point [1]",
                        (DocumentInline("First reported point "), citation),
                    ),
                ),
            ),
            DocumentBlock(
                BlockKind.TABLE,
                "Date\tReported event",
                table=DocumentTable(
                    "Table 1",
                    ("Date", "Reported event"),
                    (
                        (
                            DocumentTableCell("2026-09-12"),
                            DocumentTableCell(
                                "Event recorded [1]",
                                (DocumentInline("Event recorded "), citation),
                            ),
                        ),
                    ),
                    "Chronology from retained reporting.",
                ),
            ),
            DocumentBlock(BlockKind.FIGURE, "Figure 1", figure=figure),
            DocumentBlock(BlockKind.HEADING, "References"),
            DocumentBlock(
                BlockKind.REFERENCE,
                "[1] Example News. Source title. 2026-09-12. https://example.org/source",
            ),
        ),
        references=(
            DocumentReference(
                1,
                "E1",
                "Source title",
                "Example News",
                "2026-09-12",
                "2026-09-12T08:00:00Z",
                "https://example.org/source",
            ),
        ),
    )


def test_word_uses_native_tables_local_figures_and_safe_citation_links() -> None:
    content = render_docx(structured_document())
    word = Document(io.BytesIO(content))
    assert len(word.tables) == 1
    assert word.tables[0].cell(1, 1).text == "Event recorded [1]"
    assert "Figure 1. Locations reported" in "\n".join(p.text for p in word.paragraphs)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        document_xml = archive.read("word/document.xml").decode()
        relationships = archive.read("word/_rels/document.xml.rels").decode()
        assert 'w:tblHeader w:val="true"' in document_xml
        assert 'w:anchor="reference-1"' in document_xml
        assert 'w:name="reference-1"' in document_xml
        assert "Two reported locations" in document_xml
        assert "https://example.org/source" in relationships
        assert any(name.startswith("word/media/") for name in archive.namelist())


def test_pdf_has_searchable_table_figure_caption_links_and_outline() -> None:
    pdf = PdfReader(io.BytesIO(render_pdf(structured_document())))
    text = "\n".join(page.extract_text() for page in pdf.pages)
    for value in (
        "A source-backed finding [1]",
        "Event recorded [1]",
        "Figure 1. Locations reported",
        "https://example.org/source",
    ):
        assert value in text
    assert pdf.outline
    assert any(page.get("/Annots") for page in pdf.pages)


def test_pdf_repeats_table_header_after_page_breaks() -> None:
    rows = tuple(
        (
            DocumentTableCell(f"2026-09-{(index % 30) + 1:02d}"),
            DocumentTableCell(f"Reported event {index} with enough detail to wrap."),
        )
        for index in range(100)
    )
    document = ReportDocument(
        "Long chronology",
        "Report fixture | version 1",
        (
            DocumentBlock(BlockKind.TITLE, "Long chronology"),
            DocumentBlock(
                BlockKind.TABLE,
                "Date\tReported event",
                table=DocumentTable("Table 1", ("Date", "Reported event"), rows),
            ),
        ),
    )
    pdf = PdfReader(io.BytesIO(render_pdf(document)))
    assert len(pdf.pages) > 1
    assert all("Date" in (page.extract_text() or "") for page in pdf.pages)


def test_html_uses_fixed_structure_and_embeds_only_verified_image_bytes() -> None:
    html = render_html(structured_document()).decode()
    assert "<table>" in html and "<thead>" in html and "<figure>" in html
    assert 'href="#reference-1"' in html
    assert 'href="https://example.org/source"' in html
    assert 'src="data:image/png;base64,' in html
    assert "Two reported locations on a blue background." in html


def test_invalid_or_excessive_figure_data_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    with pytest.raises(InvalidRequest, match="invalid figure"):
        render_pdf(structured_document(figure_content=b"not an image"))
    monkeypatch.setattr(figure_validation, "MAX_DOCUMENT_FIGURE_BYTES", 1)
    with pytest.raises(InvalidRequest, match="byte limit"):
        render_docx(structured_document())


def test_non_http_reference_never_creates_a_live_link() -> None:
    document = structured_document()
    unsafe = DocumentReference(
        1, "E1", "Source title", "Example News", None, "2026-09-12", "file:///private"
    )
    blocks = (*document.blocks[:-1], DocumentBlock(BlockKind.REFERENCE, "[1] file:///private"))
    content = render_docx(
        ReportDocument(document.title, document.reference, blocks, references=(unsafe,))
    )
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        relationships = archive.read("word/_rels/document.xml.rels").decode()
        assert "file:///private" not in relationships
