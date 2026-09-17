"""The exported product opens, paginates and reads like a considered document."""

import io
import zipfile

from docx import Document
from pypdf import PdfReader

from ase.adapters.reports.document_sections import leading_identity
from ase.adapters.reports.pdf import render_pdf
from ase.adapters.reports.pdf_content import _column_widths
from ase.adapters.reports.word import render_docx
from ase.application.reports.document import build_document
from ase.application.reports.render import render_markdown
from ase.domain.report_documents import (
    BlockKind,
    DocumentBlock,
    DocumentTable,
    DocumentTableCell,
    ReportDocument,
)
from report_documents_helpers import document_records


def _document() -> ReportDocument:
    record, version = document_records()
    return build_document(record, version)


def test_leading_title_and_metadata_form_one_masthead() -> None:
    document = _document()
    head, remaining = leading_identity(document)
    assert [block.kind for block in head] == [
        BlockKind.TITLE,
        BlockKind.METADATA,
        BlockKind.METADATA,
    ]
    assert remaining and remaining[0].kind is not BlockKind.TITLE
    assert len(head) + len(remaining) == len(document.blocks)


def test_pdf_repeats_the_same_bytes_for_the_same_frozen_version() -> None:
    document = _document()
    assert render_pdf(document) == render_pdf(document)
    assert render_docx(document) == render_docx(document)


def test_pdf_carries_a_masthead_running_head_and_numbered_pages() -> None:
    document = _document()
    pdf = PdfReader(io.BytesIO(render_pdf(document)))
    pages = [page.extract_text() or "" for page in pdf.pages]
    assert len(pages) > 1
    assert "INTELLIGENCE PRODUCT" in pages[0]
    # A numbered footer on every page, and a running head carrying the title after it.
    assert all(f"Page {number} of {len(pages)}" in page for number, page in enumerate(pages, 1))
    assert all(document.title in page for page in pages[1:])


def test_pdf_columns_share_width_by_content_rather_than_evenly() -> None:
    table = DocumentTable(
        "Chronology",
        ("Date", "Reported item"),
        (
            (
                DocumentTableCell("2026-09-05"),
                DocumentTableCell("A long reported item that needs room to wrap sensibly."),
            ),
        ),
    )
    date_width, item_width = _column_widths(table)
    assert item_width > date_width
    assert date_width >= 54.0


def test_a_long_table_still_starts_on_the_page_its_heading_opens() -> None:
    rows = tuple(
        (
            DocumentTableCell(f"2026-09-{(index % 30) + 1:02d}"),
            DocumentTableCell(f"Reported event {index} with enough detail to wrap."),
        )
        for index in range(100)
    )
    document = ReportDocument(
        "Long chronology",
        "Report fixture | version 1",
        (
            DocumentBlock(BlockKind.TITLE, "Long chronology"),
            DocumentBlock(
                BlockKind.TABLE,
                "Date\tReported event",
                table=DocumentTable("Table 1", ("Date", "Reported event"), rows),
            ),
        ),
    )
    pdf = PdfReader(io.BytesIO(render_pdf(document)))
    assert "Table 1" in (pdf.pages[0].extract_text() or "")
    assert "2026-09-01" in (pdf.pages[0].extract_text() or "")


def test_word_opens_with_the_masthead_and_repeats_head_and_page_fields() -> None:
    document = _document()
    content = render_docx(document)
    word = Document(io.BytesIO(content))
    assert word.paragraphs[0].text == "INTELLIGENCE PRODUCT"
    assert word.paragraphs[1].style.name == "Title"
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        footers = "\n".join(
            archive.read(name).decode()
            for name in archive.namelist()
            if name.startswith("word/footer")
        )
        headers = "\n".join(
            archive.read(name).decode()
            for name in archive.namelist()
            if name.startswith("word/header")
        )
    assert 'w:instr="PAGE"' in footers and 'w:instr="NUMPAGES"' in footers
    assert document.title in headers


def test_word_bands_a_review_notice_so_it_reads_as_a_notice() -> None:
    document = ReportDocument(
        "Banded",
        "Report fixture | version 1",
        (
            DocumentBlock(BlockKind.TITLE, "Banded"),
            DocumentBlock(BlockKind.WARNING, "NEEDS REVIEW: unresolved checks."),
        ),
    )
    with zipfile.ZipFile(io.BytesIO(render_docx(document))) as archive:
        body = archive.read("word/document.xml").decode()
    assert 'w:fill="FDF4E2"' in body
    assert 'w:color="CF9736"' in body


def test_markdown_is_portable_with_one_blank_line_between_blocks() -> None:
    record, version = document_records()
    markdown = render_markdown(
        record.header, version.body, version.evidence, version.quality, version.findings
    )
    assert markdown.endswith("\n") and not markdown.endswith("\n\n")
    assert "\n\n\n" not in markdown
    assert not any(line != line.rstrip() for line in markdown.splitlines())
    # Headings and the paragraphs under them must not run together when rendered.
    for heading in ("### Ground activity", "### Trajectory"):
        assert f"{heading}\n\n" in markdown
    assert "**Original title:** " in markdown
    assert "**Recorded metadata**\n\n- Published:" in markdown
    # A judgement's likelihood and confidence read as their own lines, not as a run-on
    # continuation of the statement they belong to.
    assert "\n\n  Probability: " in markdown
    assert "Watch condition: elevated.\n\n- " in markdown
