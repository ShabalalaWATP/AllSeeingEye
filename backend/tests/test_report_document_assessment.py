"""Canonical reader and exports show frozen doctrine dimensions without regrading history."""

import io
from dataclasses import replace

import pytest
from docx import Document
from pypdf import PdfReader

from ase.adapters.reports.documents import ReportDocumentRenderer
from ase.application.reports.document import DocumentBuilder, build_document
from ase.application.reports.publication_markdown import render_document_markdown
from ase.domain import judgement_assessment, source_ratings
from ase.domain.doctrine import Confidence, Probability
from ase.domain.errors import InvalidRequest
from ase.domain.events import Reliability
from ase.domain.judgement_assessment import build_report_assessment
from ase.domain.report_documents import BlockKind, DocumentTable, DocumentTableCell, ExportFormat
from ase.domain.reports import ReportingItem, ReportingTheme
from ase.domain.source_ratings import unassessed_source_rating
from report_documents_helpers import document_records


def assessed_records():
    record, version = document_records()
    evidence = (
        replace(
            version.evidence[0],
            grade="B2",
            reliability="B",
            credibility=2,
            grade_rationale="Saved first item rationale.",
            source_rating=replace(
                unassessed_source_rating(),
                status="editorial",
                assessed_grade=Reliability.B,
                publisher_reliability_assessed=True,
                basis="Saved editorial source basis.",
            ),
        ),
        replace(
            version.evidence[1],
            grade="F6",
            reliability="F",
            credibility=6,
            grade_rationale="Saved unassessed item rationale.",
        ),
        *version.evidence[2:],
    )
    judgement = replace(
        version.body.key_judgements[0],
        probability=Probability.LIKELY,
        confidence=Confidence.MODERATE,
        confidence_statement="Engine confidence ceiling: moderate. Cited support: 2 item(s). "
        "Model rationale (unverified): Independent evidence remains incomplete.",
        supporting_evidence=("E1",),
        contradicting_evidence=("E2",),
    )
    body = replace(
        version.body,
        key_judgements=(judgement,),
        reporting=(
            ReportingTheme(
                "Current reporting",
                (ReportingItem("An attributed claim.", ("E1", "E2"), "B2, F6"),),
            ),
        ),
    )
    assessment = build_report_assessment(body, evidence, ())
    assessment = replace(
        assessment,
        method_version="historical-policy-not-for-reader",
        judgements=(
            replace(
                assessment.judgements[0],
                confidence_ceiling=Confidence.HIGH,
                explanation=("A saved judgement explanation, unchanged.",),
            ),
        ),
        limitations=("A saved method limitation, unchanged.",),
    )
    return record, replace(version, body=body, evidence=evidence, assessment=assessment)


def test_exact_saved_values_and_reference_numbers_without_current_regrading(monkeypatch):
    record, version = assessed_records()

    def forbidden(*args, **kwargs):
        raise AssertionError("Historical report regrading is forbidden")

    monkeypatch.setattr(judgement_assessment, "build_report_assessment", forbidden)
    monkeypatch.setattr(source_ratings, "source_rating_for", forbidden)
    document = build_document(record, version)
    text = "\n".join(block.text for block in document.blocks)
    # The judgement line is plain words; the yardstick and its bands live in the key.
    assert "Likely, with moderate confidence." in text
    assert "PHIA" not in text.split("Key to the terms")[0]
    assert "Likely: about 55 to about 75 percent." in text.split("Key to the terms")[1]
    assert "Recorded evidence confidence limit: high." in text
    assert "Recorded final analytical confidence: moderate." in text
    assert "A saved judgement explanation, unchanged." in text
    assert "A saved method limitation, unchanged." in text
    # Grades no longer interrupt the narrative; the table at the end still records them.
    assert "Recorded source grade(s)" not in text
    assert "Source grades" in text and "UK MOD JDP 2-00" in text
    assert "F and 6" in text and "not that the reporting was false" in text
    assert "Engine confidence ceiling:" not in text and "Cited support: 2 item(s)" not in text
    assert "Independent evidence remains incomplete." in text
    assert version.body.key_judgements[0].confidence_statement.startswith("Engine confidence")
    assert [reference.evidence_label for reference in document.references] == ["E1", "E2", "E3"]
    source_table = next(
        block.table
        for block in document.blocks
        if block.table and block.table.title == "Recorded source grades"
    )
    assert source_table.rows[0][0].inlines[-1].citation_numbers == (1,)
    assert "B: Usually reliable" in source_table.rows[0][1].text
    assert source_table.rows[0][2].text == "2: Probably true"
    assert "Saved editorial source basis." in text
    assert "Saved first item rationale." in text
    assert "historical-policy-not-for-reader" not in text


def test_missing_historical_assessment_remains_missing_and_preserves_existing_item_grades():
    record, version = document_records()
    version = replace(
        version, evidence=tuple(replace(item, source_rating=None) for item in version.evidence)
    )
    document = build_document(record, version)
    text = "\n".join(block.text for block in document.blocks)
    assert "A detailed evidence assessment was not recorded for this version." in text
    assert (
        "no confidence limit or corroboration assessment has been inferred retrospectively" in text
    )
    assert "Recorded evidence confidence limit:" not in text
    assert "Source assessment basis was not recorded in this version." in text
    assert version.assessment is None


def test_unassessed_platform_grade_is_not_presented_as_publisher_reliability():
    record, version = assessed_records()
    platform = replace(
        unassessed_source_rating(),
        provenance_role="platform",
        basis="The retained B feed grade does not assess the post author.",
    )
    version = replace(
        version,
        evidence=(replace(version.evidence[0], source_rating=platform), *version.evidence[1:]),
    )
    document = build_document(record, version)
    table = next(
        block.table
        for block in document.blocks
        if block.table and block.table.title == "Recorded source grades"
    )
    assert table.rows[0][1].text == "Retained B feed grade; publisher reliability unassessed."
    assert table.rows[0][2].text == "2: Probably true"
    assert "does not assess the post author" in " ".join(block.text for block in document.blocks)


def test_missing_saved_judgement_assessment_and_missing_reference_do_not_create_new_claims():
    record, version = assessed_records()
    assessment = replace(version.assessment, judgements=())
    judgement = replace(
        version.body.key_judgements[0], supporting_evidence=("missing",), contradicting_evidence=()
    )
    body = replace(
        version.body,
        key_judgements=(judgement,),
        reporting=(),
        assessment=(),
        alternative_hypotheses=(),
    )
    document = build_document(
        record, replace(version, assessment=assessment, body=body, advocacy=None)
    )
    text = " ".join(block.text for block in document.blocks)
    assert "A detailed evidence assessment was not recorded for this judgement." in text
    assert judgement.confidence_statement in text
    assert "Engine confidence ceiling: moderate." in text
    assert "No cited source items are available" in text
    assert not document.references
    assert "Recorded evidence confidence limit:" not in text


def test_no_saved_assessment_means_existing_confidence_statement_is_preserved():
    record, version = assessed_records()
    document = build_document(record, replace(version, assessment=None))
    text = " ".join(block.text for block in document.blocks)
    assert version.body.key_judgements[0].confidence_statement in text


def test_word_pdf_and_markdown_share_source_assessments_and_separate_confidence():
    record, version = assessed_records()
    document = build_document(record, version)
    renderer = ReportDocumentRenderer()
    word = Document(io.BytesIO(renderer.render(document, ExportFormat.DOCX)))
    word_text = " ".join(
        [
            *(p.text for p in word.paragraphs),
            *(cell.text for table in word.tables for row in table.rows for cell in row.cells),
        ]
    )
    pdf = PdfReader(io.BytesIO(renderer.render(document, ExportFormat.PDF)))
    for output in (
        render_document_markdown(document),
        word_text,
        " ".join(page.extract_text() for page in pdf.pages),
    ):
        text = " ".join(output.split())
        for expected in (
            "How the assessment was made",
            "Source grades",
            "Professional Head of Intelligence Assessment (PHIA) probability yardstick",
            "Likely, with moderate confidence",
            "Recorded source grades",
            "2: Probably true",
            "6: Cannot be judged",
            "Saved editorial source basis.",
            "A saved judgement explanation, unchanged.",
            "References",
        ):
            assert expected in text, expected
        assert "historical-policy-not-for-reader" not in text
        assert version.model not in text


def test_source_tables_are_small_and_long_recorded_rationales_stay_out_of_narrow_cells():
    record, version = assessed_records()
    evidence = tuple(
        replace(
            version.evidence[0],
            label=f"E{index}",
            event_id=f"item-{index}",
            grade_rationale="A long saved rationale. " * 50,
        )
        for index in range(1, 18)
    )
    labels = tuple(item.label for item in evidence)
    body = replace(
        version.body,
        key_judgements=(
            replace(
                version.body.key_judgements[0],
                supporting_evidence=labels,
                contradicting_evidence=(),
            ),
        ),
        reporting=(),
        assessment=(),
        alternative_hypotheses=(),
    )
    document = build_document(
        record,
        replace(version, evidence=evidence, body=body, assessment=None, advocacy=None),
        include_generated_figures=False,
    )
    tables = [
        block.table
        for block in document.blocks
        if block.table and block.table.title.startswith("Recorded source grades")
    ]
    assert [len(table.rows) for table in tables] == [8, 8, 1]
    assert all(
        "long saved rationale" not in cell.text
        for table in tables
        for row in table.rows
        for cell in row
    )
    assert any("long saved rationale" in block.text for block in document.blocks)


def test_structured_table_text_is_xml_safe_and_bounded():
    builder = DocumentBuilder()
    builder.table(DocumentTable("A table", ("Column",), ((DocumentTableCell("text\x00"),),)))
    assert builder.blocks[0].table.rows[0][0].text == "text"
    assert builder.blocks[0].kind is BlockKind.TABLE
    with pytest.raises(InvalidRequest, match="size limit"):
        builder.table(DocumentTable("Too much", ("Column",), ((DocumentTableCell("x" * 16001),),)))


def test_combined_support_and_opposition_preserve_all_citations_with_bounded_runs():
    _, version = document_records()
    evidence = tuple(replace(version.evidence[0], label=f"E{index}") for index in range(1, 41))
    builder = DocumentBuilder(evidence)
    builder.cited("Supporting and contrary evidence.", tuple(item.label for item in evidence))
    runs = builder.blocks[0].inlines
    assert all(len(run.citation_numbers) <= 32 for run in runs)
    assert [number for run in runs for number in run.citation_numbers] == list(range(1, 41))
    assert len(builder.references()) == 40
