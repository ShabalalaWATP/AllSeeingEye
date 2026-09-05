"""Document content, rendering safety, limits, comparison identity and access paths."""

from __future__ import annotations

import io
import zipfile
from dataclasses import replace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from docx import Document
from pypdf import PdfReader

from ase.adapters.reports import documents
from ase.adapters.reports.documents import ReportDocumentRenderer
from ase.application.reports.comparison import compare_versions
from ase.application.reports.document import (
    MAX_BLOCK_CHARS,
    MAX_DOCUMENT_CHARS,
    DocumentBuilder,
    build_document,
)
from ase.application.reports.exports import CompareReportsUseCase, ExportReportUseCase
from ase.domain.errors import Forbidden, InvalidRequest, RateLimited
from ase.domain.report_documents import BlockKind, ChangeKind, ExportFormat
from ase.domain.reports import ReportBody, ReportStatus
from ase.domain.users import User
from report_documents_helpers import document_records


def test_structured_exports_contain_headers_warnings_citations_and_frozen_provenance() -> None:
    record, version = document_records()
    document = build_document(record, version)
    renderer = ReportDocumentRenderer()
    word_bytes = renderer.render(document, ExportFormat.DOCX)
    word = Document(io.BytesIO(word_bytes))
    word_text = "\n".join(p.text for p in word.paragraphs)
    pdf = PdfReader(io.BytesIO(renderer.render(document, ExportFormat.PDF)))
    pdf_text = "\n".join(page.extract_text() for page in pdf.pages)
    for text in (word_text, pdf_text):
        for value in (
            record.title,
            "version 1",
            "NEEDS REVIEW",
            "Key judgements",
            "Direction",
            "PIR-1",
            "Devil's advocacy",
            "Gaps and collection",
            "Quality of information",
            "Validator findings",
            "Frozen evidence annex",
            "Captured:",
            "Content hash:",
            "[E1, E2]",
            "Reliability:",
            "Archive URL:",
            "Flags:",
            version.evidence[0].content_hash,
        ):
            assert value in text, value
    assert len(pdf.pages) >= 3
    assert word.paragraphs[0].style.name == "Title"
    assert any(p.style.name == "Heading 1" for p in word.paragraphs)
    assert "version 1" in word.sections[0].footer.paragraphs[0].text


def test_markup_is_literal_no_external_resources_and_unicode_remains_recoverable() -> None:
    record, version = document_records()
    hostile = '<img src="file:///etc/passwd"/><a href="https://private.test">秘密</a>\x00'
    version = replace(version, body=replace(version.body, sourcing_statement=hostile))
    document = build_document(record, version)
    renderer = ReportDocumentRenderer()
    word_bytes = renderer.render(document, ExportFormat.DOCX)
    word = Document(io.BytesIO(word_bytes))
    assert hostile[:-1] in "\n".join(p.text for p in word.paragraphs)
    with zipfile.ZipFile(io.BytesIO(word_bytes)) as archive:
        relationships = "\n".join(
            archive.read(name).decode() for name in archive.namelist() if name.endswith(".rels")
        )
        assert 'TargetMode="External"' not in relationships
        assert not any(name.startswith("word/media/") for name in archive.namelist())
    pdf = PdfReader(io.BytesIO(renderer.render(document, ExportFormat.PDF)))
    text = "\n".join(page.extract_text() for page in pdf.pages)
    assert '<img src="file:///etc/passwd"/>' in text
    assert "[U+79D8][U+5BC6]" in text
    assert all(not page.get("/Annots") for page in pdf.pages)


def test_old_version_keeps_its_frozen_period_and_failed_empty_documents_are_explicit() -> None:
    record, version = document_records()
    record.period_from = record.period_to
    frozen_header = next(
        line for line in version.markdown.splitlines() if line.startswith("Template:")
    )
    assert frozen_header in [b.text for b in build_document(record, version).blocks]
    failed = replace(
        version,
        body=ReportBody(),
        evidence=(),
        findings=(),
        direction=None,
        advocacy=None,
        markdown="",
        status=ReportStatus.FAILED,
    )
    document = build_document(record, failed)
    text = "\n".join(b.text for b in document.blocks)
    assert (
        "FAILED GENERATION" in text and "No key judgements" in text and "No frozen evidence" in text
    )
    assert "Template: intsum" in text
    assert ReportDocumentRenderer().render(document, ExportFormat.PDF).startswith(b"%PDF")


def test_export_bounds_and_concurrency_are_enforced(monkeypatch: pytest.MonkeyPatch) -> None:
    record, version = document_records()
    for oversized in (
        replace(version, markdown="x" * (MAX_DOCUMENT_CHARS * 2 + 1)),
        replace(
            version, body=replace(version.body, sourcing_statement="x" * (MAX_BLOCK_CHARS + 1))
        ),
        replace(version, evidence=(version.evidence[0],) * 101),
    ):
        with pytest.raises(InvalidRequest):
            build_document(record, oversized)
    builder = DocumentBuilder()
    builder.characters = MAX_DOCUMENT_CHARS
    with pytest.raises(InvalidRequest):
        builder.add("more")
    builder = DocumentBuilder()
    for _ in range(2000):
        builder.add("", BlockKind.TEXT)
    with pytest.raises(InvalidRequest):
        builder.add("")
    assert documents._RENDER_SLOTS.acquire(blocking=False)
    assert documents._RENDER_SLOTS.acquire(blocking=False)
    try:
        with pytest.raises(RateLimited):
            ReportDocumentRenderer().render(build_document(record, version), ExportFormat.PDF)
    finally:
        documents._RENDER_SLOTS.release()
        documents._RENDER_SLOTS.release()

    def fail(_: object) -> bytes:
        raise RuntimeError("failed")

    monkeypatch.setattr(documents, "render_pdf", fail)
    with pytest.raises(RuntimeError):
        ReportDocumentRenderer().render(build_document(record, version), ExportFormat.PDF)
    assert documents._RENDER_SLOTS.acquire(blocking=False)
    documents._RENDER_SLOTS.release()


def test_comparison_tracks_content_validation_and_evidence_by_event_identity() -> None:
    record, version = document_records()
    changed_body = replace(
        version.body,
        sourcing_statement="Updated",
        collection_recommendations=("New recommendation", "Added"),
        assumptions=(),
    )
    revised = replace(
        version,
        number=2,
        body=changed_body,
        direction=None,
        advocacy=None,
        status=ReportStatus.READY,
        findings=(),
        evidence=(
            replace(version.evidence[0], grade="D4", label="E2"),
            replace(version.evidence[1], event_id="replacement", label="E1"),
        ),
    )
    comparison = compare_versions(record, version, revised)
    assert comparison == compare_versions(record, version, revised)
    assert comparison.from_version == 1 and comparison.to_version == 2
    assert {change.kind for change in comparison.changes} == set(ChangeKind)
    assert any(
        change.path.endswith(".grade")
        and change.after == "D4"
        and change.kind == ChangeKind.CHANGED
        for change in comparison.changes
    )
    assert any(
        change.section == "Evidence"
        and "replacement" in change.path
        and change.kind == ChangeKind.ADDED
        for change in comparison.changes
    )
    assert any(
        change.section == "Evidence"
        and version.evidence[1].event_id in change.path
        and change.kind == ChangeKind.REMOVED
        for change in comparison.changes
    )
    assert any(change.section == "Validation" for change in comparison.changes)
    assert compare_versions(record, version, version).changes == ()
    assert any(
        change.kind == ChangeKind.ADDED
        for change in compare_versions(record, revised, version).changes
    )


async def test_use_cases_reuse_authorised_reader_and_reject_bad_version_numbers(user: User) -> None:
    record, version = document_records()
    reader = AsyncMock()
    reader.execute.return_value = (record, version)
    exporter = ExportReportUseCase(reader, ReportDocumentRenderer())
    result = await exporter.execute(user, record.id, ExportFormat.DOCX, 1)
    assert result.filename == f"report-{record.id}-v1.docx"
    assert result.content.startswith(b"PK")
    reader.execute.assert_awaited_with(user, record.id, 1)
    comparator = CompareReportsUseCase(reader)
    assert (await comparator.execute(user, record.id, 1, 1)).changes == ()
    with pytest.raises(InvalidRequest):
        await exporter.execute(user, record.id, ExportFormat.PDF, 0)
    with pytest.raises(InvalidRequest):
        await comparator.execute(user, record.id, 0, 1)
    reader.execute.side_effect = Forbidden()
    with pytest.raises(Forbidden):
        await exporter.execute(user, uuid4(), ExportFormat.PDF)
    with pytest.raises(Forbidden):
        await comparator.execute(user, uuid4(), 1, 2)
