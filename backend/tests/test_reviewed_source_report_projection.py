"""Explicit reviewer snapshots reach one report reader and every canonical export."""

import io

from docx import Document
from pypdf import PdfReader

from helpers import USER_PASSWORD, bearer, create_user, login_token
from source_projection_helpers import BODY
from source_review_test_helpers import review_payload, seed_reviews


async def _snapshot(client, path, headers):
    reliability = await client.post(
        path + "/source-reviews", headers=headers, json=review_payload()
    )
    assert reliability.status_code == 201, reliability.text
    credibility = await client.post(
        path + "/source-reviews",
        headers=headers,
        json=review_payload(
            kind="credibility", reliability=None, expertise_basis=None, credibility=1
        ),
    )
    assert credibility.status_code == 201, credibility.text
    frozen = await client.post(
        path + "/source-assessment-snapshots",
        headers=headers,
        json={"subjects": {BODY.key_judgements[0].id: "population-statistics"}},
    )
    assert frozen.status_code == 201, frozen.text
    return frozen.json()


async def test_selected_snapshot_projects_to_reader_markdown_word_and_pdf_without_regrading(
    app, container, client, user
):
    record, version, path = await seed_reviews(app, container, user)
    headers = bearer(await login_token(client, user.email, USER_PASSWORD))
    snapshot = await _snapshot(client, path, headers)
    report_path = f"/api/reports/{record.id}"
    query = {"version": 1, "source_snapshot_id": snapshot["id"]}

    original = await client.get(report_path, headers=headers, params={"version": 1})
    selected = await client.get(report_path, headers=headers, params=query)
    assert original.status_code == selected.status_code == 200
    assert original.json()["version"]["reviewed_source_snapshot"] is None
    selected_version = selected.json()["version"]
    assert selected_version["reviewed_source_snapshot"]["id"] == snapshot["id"]
    assert selected_version["body"] == original.json()["version"]["body"]
    assert selected_version["status"] == original.json()["version"]["status"]
    publication = " ".join(block["text"] for block in selected_version["publication"]["blocks"])
    original_publication = " ".join(
        block["text"] for block in original.json()["version"]["publication"]["blocks"]
    )
    assert "Reviewed source assessment snapshot" in publication
    assert "A: Completely reliable" in publication
    assert "1: Confirmed by other sources" in publication
    assert "Reviewer checked the published methodology" in publication
    assert "Reviewed source assessment snapshot" not in original_publication
    assert selected_version["evidence"] == original.json()["version"]["evidence"]
    assert version.source_assessment is not None

    markdown = await client.get(report_path + "/markdown", headers=headers, params=query)
    assert markdown.status_code == 200, markdown.text
    assert "Reviewed source assessment snapshot" in markdown.text
    assert "A: Completely reliable" in markdown.text
    assert snapshot["id"] in markdown.text
    word = await client.get(report_path + "/export/docx", headers=headers, params=query)
    assert word.status_code == 200, word.text
    docx = Document(io.BytesIO(word.content))
    word_text = " ".join(
        [
            *(paragraph.text for paragraph in docx.paragraphs),
            *(cell.text for table in docx.tables for row in table.rows for cell in row.cells),
        ]
    )
    assert "Reviewed source assessment snapshot" in word_text
    assert "A: Completely reliable" in word_text
    pdf = await client.get(report_path + "/export/pdf", headers=headers, params=query)
    assert pdf.status_code == 200, pdf.text
    # Normalise whitespace: a five-column table wraps on A4, and this asserts that the
    # export carries the reviewed wording, not that it lands on a single line.
    pdf_text = " ".join(
        " ".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf.content)).pages).split()
    )
    assert "Reviewed source assessment snapshot" in pdf_text
    assert "A: Completely reliable" in pdf_text


async def test_snapshot_requires_explicit_exact_version_and_current_authority(
    app, container, client, user
):
    record, _, path = await seed_reviews(app, container, user)
    headers = bearer(await login_token(client, user.email, USER_PASSWORD))
    snapshot = await _snapshot(client, path, headers)
    report_path = f"/api/reports/{record.id}"
    for suffix in ("", "/markdown", "/export/docx"):
        missing_version = await client.get(
            report_path + suffix, headers=headers, params={"source_snapshot_id": snapshot["id"]}
        )
        assert missing_version.status_code == 422
        wrong_version = await client.get(
            report_path + suffix,
            headers=headers,
            params={"version": 2, "source_snapshot_id": snapshot["id"]},
        )
        assert wrong_version.status_code == 404
    outsider = await create_user(
        container, email="source-projection-outsider@example.com", password=USER_PASSWORD
    )
    outsider_headers = bearer(await login_token(client, outsider.email, USER_PASSWORD))
    denied = await client.get(
        report_path,
        headers=outsider_headers,
        params={"version": 1, "source_snapshot_id": snapshot["id"]},
    )
    assert denied.status_code == 404


async def test_new_review_never_rewrites_selected_or_original_report_projection(
    app, container, client, user
):
    record, _, path = await seed_reviews(app, container, user)
    headers = bearer(await login_token(client, user.email, USER_PASSWORD))
    first = await client.post(path + "/source-reviews", headers=headers, json=review_payload())
    assert first.status_code == 201, first.text
    subjects = {BODY.key_judgements[0].id: "population-statistics"}
    first_snapshot = await client.post(
        path + "/source-assessment-snapshots",
        headers=headers,
        json={"subjects": subjects},
    )
    assert first_snapshot.status_code == 201, first_snapshot.text
    report_path = f"/api/reports/{record.id}"
    original_markdown = await client.get(
        report_path + "/markdown", headers=headers, params={"version": 1}
    )
    first_markdown = await client.get(
        report_path + "/markdown",
        headers=headers,
        params={"version": 1, "source_snapshot_id": first_snapshot.json()["id"]},
    )
    assert original_markdown.status_code == first_markdown.status_code == 200
    correction = await client.post(
        path + "/source-reviews",
        headers=headers,
        json=review_payload(
            previous_id=first.json()["review"]["id"],
            reliability="F",
            expertise_basis=None,
            basis="Review could not establish competence for this subject.",
        ),
    )
    assert correction.status_code == 201, correction.text
    second_snapshot = await client.post(
        path + "/source-assessment-snapshots",
        headers=headers,
        json={"subjects": subjects},
    )
    assert second_snapshot.status_code == 201, second_snapshot.text
    original = await client.get(report_path, headers=headers, params={"version": 1})
    first_selected = await client.get(
        report_path,
        headers=headers,
        params={"version": 1, "source_snapshot_id": first_snapshot.json()["id"]},
    )
    second_selected = await client.get(
        report_path,
        headers=headers,
        params={"version": 1, "source_snapshot_id": second_snapshot.json()["id"]},
    )
    assert original.status_code == first_selected.status_code == second_selected.status_code == 200
    assert original.json()["version"]["reviewed_source_snapshot"] is None
    first_text = " ".join(
        block["text"] for block in first_selected.json()["version"]["publication"]["blocks"]
    )
    second_text = " ".join(
        block["text"] for block in second_selected.json()["version"]["publication"]["blocks"]
    )
    assert "A: Completely reliable" in first_text
    assert "F: Cannot be judged" in second_text
    assert "A: Completely reliable" not in second_text
    assert "Reviewed source assessment snapshot" not in " ".join(
        block["text"] for block in original.json()["version"]["publication"]["blocks"]
    )
    unchanged_original_markdown = await client.get(
        report_path + "/markdown", headers=headers, params={"version": 1}
    )
    unchanged_first_markdown = await client.get(
        report_path + "/markdown",
        headers=headers,
        params={"version": 1, "source_snapshot_id": first_snapshot.json()["id"]},
    )
    assert unchanged_original_markdown.content == original_markdown.content
    assert unchanged_first_markdown.content == first_markdown.content
