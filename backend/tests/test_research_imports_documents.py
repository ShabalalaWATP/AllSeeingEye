"""Synthetic local PDF and DOCX extraction; no operator documents or network calls."""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, NameObject
from reportlab.pdfgen.canvas import Canvas

from ase.adapters.research_imports import extract_upload
from ase.adapters.research_imports.models import MAX_XML_BYTES, ImportRejected


def pdf_bytes(*texts: str) -> bytes:
    output = io.BytesIO()
    canvas = Canvas(output)
    for text in texts:
        canvas.drawString(72, 720, text)
        canvas.showPage()
    canvas.save()
    return output.getvalue()


def docx_zip(document: bytes, extra: dict[str, bytes] | None = None) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr("word/document.xml", document)
        for name, contents in (extra or {}).items():
            archive.writestr(name, contents)
    return output.getvalue()


def test_pdf_page_locators_text_and_omissions_are_explicit() -> None:
    result = extract_upload(pdf_bytes("First finding", "", "Third finding"), "research.pdf")
    assert [unit.reference for unit in result.units] == ["PDF page 1", "PDF page 3"]
    assert "First finding" in result.units[0].text
    assert any("1 page(s)" in limit for limit in result.limitations)
    assert any("physical page numbers" in limit for limit in result.limitations)


def test_pdf_encryption_and_page_limits_are_rejected() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("synthetic-password")
    stream = io.BytesIO()
    writer.write(stream)
    with pytest.raises(ImportRejected, match="Encrypted"):
        extract_upload(stream.getvalue(), "secret.pdf")
    with pytest.raises(ImportRejected, match="page limit"):
        extract_upload(pdf_bytes(*(["x"] * 51)), "large.pdf")


def test_pdf_decoded_stream_bomb_is_rejected_before_text_extraction() -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=100, height=100)
    content = DecodedStreamObject()
    content.set_data(b" " * (MAX_XML_BYTES + 1))
    page[NameObject("/Contents")] = content.flate_encode()
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ImportRejected):
        extract_upload(output.getvalue(), "compressed.pdf")


def test_docx_body_order_cites_table_paragraphs_without_inventing_pages() -> None:
    document = Document()
    document.add_paragraph("Source statement")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Claim"
    document.add_paragraph("Conclusion")
    document.sections[0].header.paragraphs[0].text = "Header is omitted"
    output = io.BytesIO()
    document.save(output)
    result = extract_upload(output.getvalue(), "report.docx")
    assert [unit.text for unit in result.units] == [
        "Source statement",
        "Name",
        "Claim",
        "Conclusion",
    ]
    assert [unit.reference for unit in result.units] == [
        f"DOCX body paragraph {i}" for i in range(1, 5)
    ]
    assert any("page numbers are unavailable" in limit for limit in result.limitations)


@pytest.mark.parametrize(
    ("source", "filename"),
    [(b"not PDF", "x.pdf"), (b"%PDF-broken SECRET", "x.pdf"), (b"not ZIP", "x.docx")],
)
def test_binary_parse_failures_do_not_echo_document_content(source: bytes, filename: str) -> None:
    with pytest.raises(ImportRejected) as error:
        extract_upload(source, filename)
    assert "SECRET" not in str(error.value)


@pytest.mark.parametrize("entry", ["../outside", "/absolute", "x\\y", "C:drive"])
def test_docx_rejects_unsafe_archive_paths(entry: str) -> None:
    source = docx_zip(b"<document/>", {entry: b"x"})
    if entry == "x\\y":
        # Windows ZipInfo normalises separators when constructing a fixture; put
        # the hostile spelling back into both local and central-directory names.
        source = source.replace(b"x/y", b"x\\y")
    with pytest.raises(ImportRejected, match="unsafe archive"):
        extract_upload(source, "bad.docx")


def test_docx_rejects_expansion_bombs_macros_and_external_entities() -> None:
    with pytest.raises(ImportRejected, match="expansion"):
        extract_upload(docx_zip(b"x" * (MAX_XML_BYTES + 1)), "bomb.docx")
    with pytest.raises(ImportRejected, match="Macro-enabled"):
        extract_upload(docx_zip(b"<document/>", {"word/vbaProject.bin": b"x"}), "macro.docx")
    xml = b'<!DOCTYPE x [<!ENTITY xxe SYSTEM "file:///secret">]><x>&xxe;</x>'
    with pytest.raises(ImportRejected, match="malformed"):
        extract_upload(docx_zip(xml), "entity.docx")


def test_docx_requires_valid_document_structure() -> None:
    with pytest.raises(ImportRejected, match="Word document body"):
        extract_upload(docx_zip(b"<document/>"), "wrong.docx")
    xml = (
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b"<w:body>" + b"<w:p>" * 34 + b"<w:t>x</w:t>" + b"</w:p>" * 34 + b"</w:body></w:document>"
    )
    with pytest.raises(ImportRejected, match="structure limit"):
        extract_upload(docx_zip(xml), "deep.docx")


def test_docx_tab_and_break_are_retained() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("A\tB\nC")
    output = io.BytesIO()
    document.save(output)
    assert extract_upload(output.getvalue(), "notes.docx").units[0].text == "A\tB\nC"


def test_docx_rejects_archive_directory_budgets_before_xml_parsing() -> None:
    with pytest.raises(ImportRejected, match="archive directory"):
        extract_upload(b"PK\x03\x04truncated", "bad.docx")
    entries = {f"word/extra-{number}": b"x" for number in range(257)}
    with pytest.raises(ImportRejected, match="archive directory"):
        extract_upload(docx_zip(b"<document/>", entries), "many.docx")
